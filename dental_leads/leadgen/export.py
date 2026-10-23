"""Write leads.csv and leads.xlsx."""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("score", "Score"),
    ("name", "Practice Name"),
    ("city", "City"),
    ("head_dentist", "Head Dentist"),
    ("email", "Email"),
    ("email_source", "Email Source"),
    ("email_status", "Email Status"),
    ("phone", "Phone"),
    ("website", "Website"),
    ("rating", "Rating"),
    ("review_count", "Review Count"),
    ("age_display", "Est. Practice Age"),
    ("velocity_display", "Review Velocity"),
    ("why", "Why This Lead"),
    ("maps_url", "Google Maps Link"),
    # Extra detail
    ("address", "Address"),
    ("practice_types_display", "Practice Type"),
    ("dentist_count", "Dentists Found"),
    ("other_dentists", "Other Dentists Found"),
    ("other_emails", "Other Emails"),
    ("age_basis", "Age Basis"),
    ("domain_created", "Domain Registered"),
    ("established_year", "Website Est. Year"),
    ("oldest_review", "Oldest Review Seen"),
    ("rating_change", "Rating Change"),
    ("tx_license", "TX License (TSBDE)"),
    ("business_status", "Business Status"),
    ("exclude_reason", "Exclude Reason"),
    ("place_id", "Place ID"),
]
WIDTHS = {"Practice Name": 34, "Head Dentist": 22, "Email": 30, "Email Status": 34, "Website": 34,
          "Why This Lead": 80, "Google Maps Link": 30, "Address": 40, "Age Basis": 40,
          "Other Emails": 30, "Other Dentists Found": 30}


def to_frame(leads):
    df = pd.DataFrame([{label: lead.get(key) for key, label in COLUMNS} for lead in leads])
    if df.empty:
        return pd.DataFrame(columns=[label for _, label in COLUMNS])
    return df.sort_values(["Score", "Review Count"], ascending=[False, False])


def write_outputs(leads, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    keep = [l for l in leads if not l["exclude_reason"]]
    excluded = [l for l in leads if l["exclude_reason"]]
    df, df_x = to_frame(keep), to_frame(excluded)

    csv_path, xlsx_path = out_dir / "leads.csv", out_dir / "leads.xlsx"
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Leads", index=False)
        df_x.to_excel(xw, sheet_name="Excluded", index=False)
        for name in ("Leads", "Excluded"):
            _format(xw.book[name])
    return csv_path, xlsx_path


def _format(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    headers = [c.value for c in ws[1]]
    for idx, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = WIDTHS.get(h, max(12, len(str(h)) + 2))
        if h in ("Website", "Google Maps Link"):
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                c = row[0]
                if c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
        if h == "Email Status":
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                c = row[0]
                if c.value and c.value.startswith("pattern guess"):
                    c.font = Font(color="C00000", italic=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
