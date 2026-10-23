"""Full pipeline against local fake servers: Places API, RDAP and one practice website."""
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import pandas as pd
import pytest

from leadgen import age, places, pipeline
from leadgen.config import load_config


def _place(pid, name, website=None, rating=4.9, count=40, status="OPERATIONAL", city="Plano", reviews=()):
    p = {"id": pid, "displayName": {"text": name}, "formattedAddress": f"1 Main St, {city}, TX 75024, USA",
         "addressComponents": [{"longText": city, "shortText": city, "types": ["locality", "political"]}],
         "nationalPhoneNumber": "(972) 555-0100", "rating": rating, "userRatingCount": count,
         "businessStatus": status, "types": ["dentist", "health"], "googleMapsUri": f"https://maps.google.com/?cid={pid}",
         "reviews": [{"publishTime": d} for d in reviews]}
    if website:
        p["websiteUri"] = website
    return p


SITE = {
    "/robots.txt": "User-agent: *\nDisallow: /private\n",
    "/": """<html><head><title>Bright Creek Dental | Plano Family Dentist</title></head><body>
        <a href="/meet-dr-smith">Meet Dr. Smith</a> <a href="/contact">Contact</a> <a href="/private/x">Team</a>
        <h1>Family &amp; cosmetic dentistry in Plano</h1><p>Proudly serving Plano since 2024. Cosmetic veneers.</p>
        </body></html>""",
    "/meet-dr-smith": """<html><body><h1>Meet Dr. Jane Smith, DDS</h1>
        <p>Dr. Jane Smith, DDS is the founder of Bright Creek Dental and opened the practice in 2024.</p>
        <p>Dr. Smith loves cosmetic dentistry.</p></body></html>""",
    "/contact": """<html><body><a href="mailto:hello@brightcreek.test">hello@brightcreek.test</a>
        <a href="mailto:careers@brightcreek.test">jobs</a></body></html>""",
    "/private/x": "<html><body>Dr. Secret Person, DDS</body></html>",
}


class Fake:
    def __init__(self):
        self.places_calls = 0
        self.site_paths = []


@pytest.fixture
def servers():
    fake = Fake()

    class SiteHandler(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            path = self.path.split("?")[0]
            fake.site_paths.append(path)
            body = SITE.get(path)
            self.send_response(200 if body else 404)
            self.send_header("Content-Type", "text/plain" if path == "/robots.txt" else "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write((body or "not found").encode())

    site = ThreadingHTTPServer(("127.0.0.1", 0), SiteHandler)
    site_url = f"http://127.0.0.1:{site.server_port}/"

    class ApiHandler(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def _send(self, obj):
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps(obj).encode())

        def do_POST(self):
            fake.places_calls += 1
            assert self.headers["X-Goog-Api-Key"] == "test-key"
            assert "places.reviews" in self.headers["X-Goog-FieldMask"]
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            if body.get("pageToken") == "page2":
                self._send({"places": [
                    _place("p3", "Aspen Dental - Plano", rating=4.2, count=300),
                    _place("p4", "Old Town Dental", status="CLOSED_PERMANENTLY"),
                ]})
            else:
                self._send({"places": [
                    _place("p1", "Bright Creek Dental", site_url, rating=4.9, count=48, reviews=["2025-02-01T10:00:00Z"]),
                    _place("p2", "Dr. Raj Patel Dentistry", None, rating=4.6, count=220,
                           reviews=["2016-05-01T10:00:00Z", "2026-01-01T00:00:00Z"]),
                    _place("p5", "Lakeside Dental", None, city="Richardson"),   # outside target cities
                ][:2] if body["pageSize"] == 2 else [], "nextPageToken": "page2"})

        def do_GET(self):  # RDAP
            self._send({"events": [{"eventAction": "registration", "eventDate": "2023-11-15T00:00:00Z"}]})

    api = ThreadingHTTPServer(("127.0.0.1", 0), ApiHandler)
    for s in (site, api):
        threading.Thread(target=s.serve_forever, daemon=True).start()
    fake.api = f"http://127.0.0.1:{api.server_port}"
    yield fake
    site.shutdown()
    api.shutdown()


def test_full_run_and_cached_rerun(servers, tmp_path, monkeypatch):
    monkeypatch.setattr(places, "SEARCH_URL", servers.api + "/v1/places:searchText")
    monkeypatch.setattr(age, "RDAP_URL", servers.api + "/rdap/{domain}")
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY", "test-key")
    monkeypatch.delenv("HUNTER_API_KEY", raising=False)
    cfg = load_config(env_path=tmp_path / "none.env")
    cfg.update(grid_cell_km=100, queries=["dentist in {city}, TX"], database=str(tmp_path / "leads.sqlite"),
               output_dir=str(tmp_path / "out"))
    cfg["cache"]["path"] = str(tmp_path / "cache.sqlite")
    cfg["places"]["page_size"] = 2
    cfg["crawl"]["delay_seconds"] = 0

    leads = pipeline.run(cfg, ["Plano"], yes=True)
    assert servers.places_calls == 2                      # page 1 + page 2
    assert "/private/x" not in servers.site_paths         # robots.txt respected
    by_id = {l["place_id"]: l for l in leads}
    assert set(by_id) == {"p1", "p2", "p3", "p4"}

    a = by_id["p1"]
    assert a["head_dentist"] == "Dr. Jane Smith" and a["dentist_count"] == 1
    assert a["email"] == "hello@brightcreek.test" and a["email_source"] == "website"
    assert a["email_status"].startswith("verified on site")
    assert a["established_year"] == 2024 and a["age_basis"] == "website says est. 2024"
    assert "cosmetic" in a["practice_types"]

    b = by_id["p2"]   # no website: name from Google listing, no email, age from oldest review
    assert b["head_dentist"] == "Dr. Raj Patel" and b["email"] == ""
    assert b["oldest_review"] == "2016-05-01"
    assert a["score"] > b["score"] > 0

    assert by_id["p3"]["exclude_reason"].startswith("corporate chain") and by_id["p3"]["score"] == 0
    assert by_id["p4"]["exclude_reason"] == "permanently closed"

    df = pd.read_csv(tmp_path / "out" / "leads.csv")
    assert list(df["Practice Name"]) == ["Bright Creek Dental", "Dr. Raj Patel Dentistry"]
    assert list(df.columns[:15]) == ["Score", "Practice Name", "City", "Head Dentist", "Email", "Email Source",
                                     "Email Status", "Phone", "Website", "Rating", "Review Count",
                                     "Est. Practice Age", "Review Velocity", "Why This Lead", "Google Maps Link"]
    wb = openpyxl.load_workbook(tmp_path / "out" / "leads.xlsx")
    assert wb.sheetnames == ["Leads", "Excluded"] and wb["Excluded"].max_row == 3

    # Rerun: everything comes from cache, so Google isn't called again.
    calls_before = servers.places_calls
    pipeline.run(cfg, ["Plano"], yes=True)
    assert servers.places_calls == calls_before


def test_limit_stops_search_early(servers, tmp_path, monkeypatch):
    monkeypatch.setattr(places, "SEARCH_URL", servers.api + "/v1/places:searchText")
    monkeypatch.setattr(age, "RDAP_URL", servers.api + "/rdap/{domain}")
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY", "test-key")
    cfg = load_config(env_path=tmp_path / "none.env")
    cfg.update(grid_cell_km=1, database=str(tmp_path / "l.sqlite"), output_dir=str(tmp_path / "out"))
    cfg["cache"]["path"] = str(tmp_path / "c.sqlite")
    cfg["places"]["page_size"] = 2
    cfg["crawl"]["delay_seconds"] = 0
    leads = pipeline.run(cfg, ["Plano"], limit=2, yes=True)
    assert len(leads) == 2 and servers.places_calls == 1
