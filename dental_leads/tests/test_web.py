"""The web app, driven through Flask's test client against the same fake servers as test_end_to_end."""
import base64
import io
import time

import openpyxl
import pytest
import yaml

from leadgen import age, config, places
from leadgen.web import create_app
from leadgen.web.settings_schema import SECTIONS, field_name, get_path
from test_end_to_end import servers  # noqa: F401  (fixture)


@pytest.fixture
def client(servers, tmp_path, monkeypatch):  # noqa: F811
    monkeypatch.setattr(places, "SEARCH_URL", servers.api + "/v1/places:searchText")
    monkeypatch.setattr(age, "RDAP_URL", servers.api + "/rdap/{domain}")
    for key in config.ENV_KEYS:   # set-then-delete so monkeypatch restores whatever save_env writes
        monkeypatch.setenv(key, "")
        monkeypatch.delenv(key)
    config.LOCAL_SETTINGS.write_text(yaml.safe_dump({
        "database": str(tmp_path / "leads.sqlite"), "output_dir": str(tmp_path / "out"),
        "cache": {"path": str(tmp_path / "cache.sqlite")}, "places": {"page_size": 2},
        "crawl": {"delay_seconds": 0},
    }))
    app = create_app()
    app.testing = True
    return app.test_client()


def _settings_form(**changes):
    """The settings form as the page would submit it, filled with current values."""
    cfg = config.load_config()
    form = {}
    for _, fields in SECTIONS:
        for path, _, kind, _ in fields:
            value = get_path(cfg, path)
            name = field_name(path)
            if kind == "bool":
                if value:
                    form[name] = "on"
            elif kind == "lines":
                form[name] = "\n".join(value)
            elif kind == "types":
                form[name] = value
            else:
                form[name] = str(value)
    form.update(changes)
    return form


def _wait(client):
    for _ in range(100):
        job = client.get("/api/job").get_json()
        if job["status"] != "running":
            return job
        time.sleep(0.1)
    raise AssertionError("job didn't finish")


def test_settings_keys_search_and_results(client, servers):  # noqa: F811
    home = client.get("/").get_data(as_text=True)
    assert "Missing" in home and "No searches yet" in home
    assert client.post("/api/run", json={"cities": ["Plano"]}).status_code == 400   # no key yet

    # Save settings: one query, big grid cell, and the Google key.
    form = _settings_form(**{"queries": "dentist in {city}, TX", "grid_cell_km": "100",
                             "GOOGLE_PLACES_API_KEY": "test-key", "action": "save"})
    resp = client.post("/settings", data=form, follow_redirects=True)
    assert "Settings saved." in resp.get_data(as_text=True)
    assert "set, ending in …-key" in resp.get_data(as_text=True)
    saved = yaml.safe_load(config.LOCAL_SETTINGS.read_text())
    assert saved["queries"] == ["dentist in {city}, TX"] and saved["places"]["page_size"] == 2  # merged, not replaced
    assert "GOOGLE_PLACES_API_KEY=test-key" in config.ENV_PATH.read_text()

    est = client.post("/api/estimate", json={"cities": ["plano"], "limit": ""}).get_json()
    assert est["has_key"] and est["max_requests"] == 3 and "Plano" in est["text"]

    job = client.post("/api/run", json={"cities": ["Plano"]}).get_json()
    assert job["status"] == "running"
    job = _wait(client)
    assert job["status"] == "done" and job["lead_count"] == 4 and job["run_id"]
    assert any("Bright Creek Dental" in line["msg"] for line in job["log"])

    page = client.get("/leads").get_data(as_text=True)
    assert "Bright Creek Dental" in page and "hello@brightcreek.test" in page and "Download Excel" in page

    csv = client.get(f"/leads/{job['run_id']}.csv")
    assert csv.status_code == 200 and b"Bright Creek Dental" in csv.data and b"Aspen" not in csv.data
    xlsx = client.get(f"/leads/{job['run_id']}.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx.data))
    assert wb.sheetnames == ["Leads", "Excluded"] and wb["Leads"].max_row == 3
    assert client.get("/leads/999.csv").status_code == 404

    home = client.get("/").get_data(as_text=True)
    assert "Ready" in home and "Finished" in home


def test_settings_validation_and_reset(client):
    form = _settings_form(**{"icp.review_count_min": "500", "icp.review_count_max": "10", "action": "save"})
    form.pop("icp.practice_types")
    page = client.post("/settings", data=form).get_data(as_text=True)
    assert "Not saved" in page and "minimum can" in page and "at least one practice type" in page
    assert yaml.safe_load(config.LOCAL_SETTINGS.read_text()).get("icp") is None

    client.post("/settings", data=_settings_form(**{"icp.min_rating": "4.0", "action": "save"}))
    assert config.load_config()["icp"]["min_rating"] == 4.0
    client.post("/settings", data={"action": "reset"})
    assert not config.LOCAL_SETTINGS.exists()
    assert config.load_config()["icp"]["min_rating"] == 4.5


def test_bad_requests_and_cross_site_posts(client):
    assert client.post("/api/estimate", json={"cities": []}).status_code == 400
    assert client.post("/api/estimate", json={"cities": ["Plano"], "limit": "abc"}).status_code == 400
    assert client.post("/api/estimate", json={"cities": ["Plano"]},
                       headers={"Origin": "https://evil.example"}).status_code == 403


def test_password_protection(servers, tmp_path):  # noqa: F811
    app = create_app(password="s3cret")
    c = app.test_client()
    assert c.get("/").status_code == 401
    good = base64.b64encode(b"me:s3cret").decode()
    assert c.get("/", headers={"Authorization": f"Basic {good}"}).status_code == 200
